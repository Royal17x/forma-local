"""Create editable PPTX slides and local PNG previews without the Codex runtime."""

import colorsys
import io
import json
import math
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


WIDTH, HEIGHT = 1280, 720
FONT_DIR = Path('C:/Windows/Fonts')


def unit(value):
    return Inches(value / 96)


def rgb(value):
    return RGBColor.from_string(value.lstrip('#'))


def derived_color(hue, saturation, lightness):
    red, green, blue = colorsys.hls_to_rgb(hue, lightness, saturation)
    return '#{:02X}{:02X}{:02X}'.format(round(red * 255), round(green * 255), round(blue * 255))


def palette_for(primary):
    if not isinstance(primary, str) or len(primary) != 7 or not primary.startswith('#'):
        primary = '#007A62'
    try:
        red, green, blue = (int(primary[index:index + 2], 16) / 255 for index in (1, 3, 5))
    except ValueError:
        red, green, blue = 0, 122 / 255, 98 / 255
    hue, _, saturation = colorsys.rgb_to_hls(red, green, blue)
    saturation = min(.72, max(.25, saturation))
    return {
        'light': '#F7F7F5', 'dark': derived_color(hue, .14, .12),
        'ink': '#202626', 'muted': '#5D6665', 'rule': '#D8DEDB',
        'accent': derived_color(hue, saturation, .32),
        'dark_accent': derived_color(hue, min(.55, saturation), .72),
        'dark_muted': '#BAC4C1',
    }


def font_at(size, bold):
    preferred = FONT_DIR / ('arialbd.ttf' if bold else 'arial.ttf')
    fallback = Path(__file__).parent / 'DejaVuSans.ttf'
    return ImageFont.truetype(str(preferred if preferred.exists() else fallback), size)


def wrap_text(draw, value, font, max_width):
    lines = []
    for paragraph in str(value or '').replace('\r', '').split('\n'):
        if not paragraph.strip():
            lines.append('')
            continue
        current = ''
        for word in paragraph.split():
            candidate = (current + ' ' + word).strip()
            if draw.textlength(candidate, font=font) <= max_width:
                current = candidate
                continue
            if current:
                lines.append(current)
                current = ''
            while word and draw.textlength(word, font=font) > max_width:
                cut = len(word) - 1
                while cut > 1 and draw.textlength(word[:cut], font=font) > max_width:
                    cut -= 1
                lines.append(word[:cut])
                word = word[cut:]
            current = word
        if current:
            lines.append(current)
    return lines or ['']


class Canvas:
    def __init__(self, prs, background):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(background)
        self.preview = Image.new('RGB', (WIDTH, HEIGHT), background)
        self.draw = ImageDraw.Draw(self.preview)

    def rect(self, x, y, width, height, fill):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, unit(x), unit(y), unit(width), unit(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        for effect in shape._element.xpath('.//a:effectRef'):
            effect.set('idx', '0')
        self.draw.rectangle((x, y, x + width, y + height), fill=fill)

    def rule(self, x, y, width, fill):
        self.rect(x, y, width, 1, fill)

    def text(self, value, x, y, width, height, size, color, bold=False):
        if value is None or value == '':
            return
        max_width = max(1, width - 12)
        for fitted_size in range(round(size), 13, -1):
            font = font_at(fitted_size, bold)
            lines = wrap_text(self.draw, value, font, max_width)
            line_height = fitted_size * 1.20
            if len(lines) * line_height <= height - 1:
                break
        else:
            fitted_size = 14
            font = font_at(fitted_size, bold)
            lines = wrap_text(self.draw, value, font, max_width)
            line_height = fitted_size * 1.20
            count = max(1, int((height - 1) // line_height))
            lines = lines[:count]
            if len(lines) == count:
                last = lines[-1]
                while last and self.draw.textlength(last + '…', font=font) > max_width:
                    last = last[:-1]
                lines[-1] = last.rstrip() + '…'

        textbox = self.slide.shapes.add_textbox(unit(x), unit(y), unit(width), unit(height))
        frame = textbox.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.0
            run = paragraph.add_run()
            run.text = line or ' '
            run.font.name = 'Arial'
            run.font.size = Pt(fitted_size * .75)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
            self.draw.text((x, y + index * line_height), line, font=font, fill=color, anchor='lt')

    def image(self, source, x, y, width, height):
        with Image.open(source) as original:
            tile = ImageOps.fit(original.convert('RGB'), (width, height), method=Image.Resampling.LANCZOS)
        buffer = io.BytesIO()
        tile.save(buffer, format='PNG')
        buffer.seek(0)
        self.slide.shapes.add_picture(buffer, unit(x), unit(y), unit(width), unit(height))
        self.preview.paste(tile, (x, y))

    def preview_text(self, value, x, y, size, color, bold=False, centered=False):
        font = font_at(size, bold)
        if centered:
            x -= self.draw.textlength(str(value), font=font) / 2
        self.draw.text((x, y), str(value), font=font, fill=color, anchor='lt')

    def notes(self, slide):
        note = str(slide.get('notes', ''))
        if slide.get('sources'):
            note += '\n\nSources:\n' + '\n'.join(slide['sources'])
        if slide.get('image'):
            note += '\nImage supplied by the user.' if slide.get('image_origin') == 'user' else \
                '\nIllustration generated locally with SDXL.'
        if slide.get('chart'):
            note += '\nChart values copied from the presentation text; native editable PowerPoint chart.'
        if slide.get('scheme'):
            note += '\nConceptual scheme generated from the slide content; editable native shapes.'
        self.slide.notes_slide.notes_text_frame.text = note


def draw_chart(canvas, chart, language, colors):
    categories = chart['categories']
    values = chart['values']
    kind = chart['type']
    data = CategoryChartData()
    data.categories = categories
    data.add_series('Value' if language == 'en' else 'Показатель', values)
    chart_kind = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE,
                  'pie': XL_CHART_TYPE.PIE}[kind]
    native = canvas.slide.shapes.add_chart(chart_kind, unit(88), unit(286), unit(1080), unit(333), data).chart
    native.has_title = False
    native.has_legend = kind == 'pie'
    if kind == 'pie':
        native.legend.position = XL_LEGEND_POSITION.RIGHT
        native.legend.include_in_layout = False
    else:
        native.category_axis.tick_labels.font.name = 'Arial'
        native.category_axis.tick_labels.font.size = Pt(12)
        native.value_axis.tick_labels.font.name = 'Arial'
        native.value_axis.tick_labels.font.size = Pt(11)
        native.value_axis.has_major_gridlines = False
    try:
        series = native.series[0]
        if kind == 'bar':
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb(colors['accent'])
        elif kind == 'line':
            series.format.line.color.rgb = rgb(colors['accent'])
            series.format.line.width = Pt(2.5)
        else:
            segments = [colors['accent'], colors['dark_accent'], '#9DB6AD', '#C7D3CE', '#566B64', '#D9E1DD', '#859B93']
            for point, color in zip(series.points, segments):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(color)
        native.plots[0].has_data_labels = True
        labels = native.plots[0].data_labels
        labels.show_value = True
        labels.position = XL_LABEL_POSITION.OUTSIDE_END if kind != 'line' else XL_LABEL_POSITION.ABOVE
        labels.font.name = 'Arial'
        labels.font.size = Pt(11)
    except (AttributeError, ValueError):
        pass

    # The PNG preview is drawn from the same numbers. The PowerPoint chart above stays editable.
    draw = canvas.draw
    accent = colors['accent']
    muted = colors['muted']
    if kind == 'pie':
        total = sum(max(0, value) for value in values)
        if total <= 0:
            return
        center = (400, 449)
        radius = 150
        box = (center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius)
        segment_colors = [accent, colors['dark_accent'], '#9DB6AD', '#C7D3CE', '#566B64', '#D9E1DD', '#859B93']
        start = -90
        for value, color in zip(values, segment_colors):
            sweep = 360 * max(0, value) / total
            draw.pieslice(box, start, start + sweep, fill=color, outline=colors['light'], width=2)
            start += sweep
        for index, (label, value) in enumerate(zip(categories, values)):
            y = 328 + index * 43
            draw.rectangle((710, y + 3, 724, y + 17), fill=segment_colors[index])
            canvas.preview_text(f'{label}  {value:g}', 739, y, 18, colors['ink'])
        return

    x0, y0, x1, y1 = 132, 320, 1147, 591
    draw.line((x0, y0, x0, y1, x1, y1), fill=colors['rule'], width=2)
    low = min(0, min(values))
    high = max(0, max(values))
    if math.isclose(high, low):
        high = low + 1
    span = high - low
    for index in range(5):
        number = low + span * index / 4
        y = y1 - (number - low) / span * (y1 - y0)
        draw.line((x0, y, x1, y), fill='#E8ECE9', width=1)
        canvas.preview_text(f'{number:g}', 82, y - 10, 14, muted)
    positions = []
    for index, (label, value) in enumerate(zip(categories, values)):
        x = x0 + (index + .5) * (x1 - x0) / len(values) if kind == 'bar' else \
            x0 + index * (x1 - x0) / max(1, len(values) - 1)
        y = y1 - (value - low) / span * (y1 - y0)
        positions.append((x, y))
        if kind == 'bar':
            bar_width = min(80, (x1 - x0) / len(values) * .54)
            zero_y = y1 - (0 - low) / span * (y1 - y0)
            draw.rectangle((x - bar_width / 2, min(y, zero_y), x + bar_width / 2, max(y, zero_y)), fill=accent)
        canvas.preview_text(f'{value:g}', x, y - 28, 15, colors['ink'], True, centered=True)
        canvas.preview_text(str(label), x, 601, 15, muted, centered=True)
    if kind == 'line':
        draw.line(positions, fill=accent, width=4, joint='curve')
        for x, y in positions:
            draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=accent)


def image_position(layout, mirror):
    if layout == 'cover':
        return 722, 0, 558, 720
    if mirror:
        return 0, 0, 480, 720
    return 760, 0, 520, 720


def render_slide(prs, slide_data, index, total, folder, context, colors):
    layout = slide_data.get('layout', 'editorial')
    dark = layout in ('cover', 'focus')
    ink = '#F6F7F4' if dark else colors['ink']
    muted = colors['dark_muted'] if dark else colors['muted']
    accent = colors['dark_accent'] if dark else colors['accent']
    canvas = Canvas(prs, colors['dark'] if dark else colors['light'])
    image_name = slide_data.get('image') if layout not in ('chart', 'scheme') else None
    with_image = bool(image_name)
    mirror = layout == 'editorial' and index % 4 == 1
    if with_image:
        image_path = (folder / image_name).resolve()
        if image_path.parent != folder.resolve():
            raise ValueError('Image must be inside project')
        canvas.image(image_path, *image_position(layout, mirror))
    content_x = 540 if mirror else 72
    content_width = (676 if mirror else 632) if with_image else 1136
    title = slide_data.get('title', '')
    summary = slide_data.get('summary', '')
    points = slide_data.get('points', [])
    language = context.get('language', 'ru')
    if layout != 'cover':
        canvas.text(f'{index + 1:02d}', content_x, 43, 42, 25, 16, accent, True)
        canvas.rule(content_x + 48, 58, content_width - 48, '#4B5552' if dark else colors['rule'])
        canvas.text(f'{index + 1:02d} / {total:02d}', content_x, 670, 100, 25, 14, muted)
    if layout == 'cover':
        canvas.rect(72, 92, 68, 5, accent)
        canvas.text(title, 72, 152, 590 if with_image else 1050, 268, 59, ink, True)
        canvas.rule(72, 463, 580 if with_image else 1050, '#4B5552')
        canvas.text(summary, 72, 490, 590 if with_image else 1050, 122, 27, muted)
        canvas.text(f'{total} ' + ('slides' if language == 'en' else 'слайдов'), 72, 658, 220, 23, 15, muted)
    elif layout == 'chart' and slide_data.get('chart'):
        chart = slide_data['chart']
        canvas.text(title, 72, 92, 1110, 84, 40, ink, True)
        canvas.text(summary, 72, 187, 1080, 61, 23, muted)
        if chart.get('unit'):
            canvas.text(('Unit: ' if language == 'en' else 'Единица: ') + chart['unit'],
                        88, 257, 260, 23, 15, muted)
        draw_chart(canvas, chart, language, colors)
        canvas.text('Data: user-provided text' if language == 'en' else 'Данные: текст пользователя',
                    72, 636, 720, 24, 14, muted)
    elif layout == 'scheme' and slide_data.get('scheme'):
        nodes = slide_data['scheme']['nodes']
        canvas.text(title, 72, 92, 1110, 84, 40, ink, True)
        canvas.text(summary, 72, 187, 1080, 62, 23, muted)
        gap = 1136 / len(nodes)
        for j, node in enumerate(nodes):
            x = 72 + j * gap
            canvas.rect(x, 329, gap - 20, 3, accent)
            canvas.text(f'{j + 1:02d}', x, 357, gap - 30, 34, 17, accent, True)
            canvas.text(node, x, 412, gap - 34, 120, 25, ink, True)
            if j < len(nodes) - 1:
                canvas.rule(x + gap - 19, 330, 19, accent)
        canvas.text('Conceptual scheme' if language == 'en' else 'Концептуальная схема',
                    72, 636, 720, 24, 14, muted)
    elif layout == 'quote':
        statement = points[0]['body'] if points else summary
        canvas.text(title, content_x, 96, content_width, 108, 38, ink, True)
        canvas.rect(content_x, 244, 54, 4, accent)
        canvas.text(statement, content_x, 275, content_width, 188, 37, ink, True)
        canvas.rule(content_x, 497, content_width, '#4B5552' if dark else colors['rule'])
        if len(points) > 1:
            canvas.text(points[1]['heading'], content_x, 515, content_width, 45, 21, accent, True)
            canvas.text(points[1]['body'], content_x, 563, content_width, 76, 19, muted)
    elif layout == 'columns' and with_image:
        canvas.text(title, 72, 96, 632, 126, 40, ink, True)
        canvas.text(summary, 72, 238, 632, 75, 22, muted)
        canvas.rule(72, 332, 632, colors['rule'])
        width = 632 / len(points)
        for j, point in enumerate(points):
            x = 72 + j * width
            if j:
                canvas.rect(x - 12, 358, 1, 262, colors['rule'])
            canvas.text(f'{j + 1:02d}', x, 356, width - 24, 25, 15, accent, True)
            canvas.text(point['heading'], x, 394, width - 28, 75, 22, ink, True)
            canvas.text(point['body'], x, 477, width - 28, 143, 19, muted)
    elif with_image:
        x = content_x
        canvas.text(title, x, 96, content_width, 126, 40, ink, True)
        canvas.text(summary, x, 238, content_width, 75, 22, muted)
        canvas.rule(x, 332, content_width, '#4B5552' if dark else colors['rule'])
        for j, point in enumerate(points):
            y = 350 + j * 149
            canvas.text(point['heading'], x, y, content_width, 49, 22, accent, True)
            canvas.text(point['body'], x, y + 53, content_width, 84, 19, ink)
            if j == 0:
                canvas.rule(x, y + 143, content_width, '#4B5552' if dark else colors['rule'])
    elif layout == 'columns':
        canvas.text(title, 72, 96, 1136, 98, 41, ink, True)
        canvas.text(summary, 72, 205, 1136, 76, 23, muted)
        canvas.rule(72, 326, 1136, colors['rule'])
        width = 1136 / len(points)
        for j, point in enumerate(points):
            x = 72 + j * width
            if j:
                canvas.rect(x - 20, 355, 1, 250, colors['rule'])
            canvas.text(f'{j + 1:02d}', x, 353, width - 42, 25, 16, accent, True)
            canvas.text(point['heading'], x, 390, width - 42, 69, 27, ink, True)
            canvas.text(point['body'], x, 470, width - 42, 142, 22, muted)
    else:
        canvas.text(title, 72, 96, 1136, 98, 41, ink, True)
        canvas.text(summary, 72, 205, 1136, 75, 23, muted)
        canvas.rule(72, 315, 1136, '#4B5552' if dark else colors['rule'])
        for j, point in enumerate(points):
            y = 354 + j * 132
            canvas.text(point['heading'], 72, y, 325, 78, 27, accent, True)
            canvas.text(point['body'], 450, y, 758, 102, 22, ink)
            if j == 0:
                canvas.rule(72, y + 124, 1136, '#4B5552' if dark else colors['rule'])
    canvas.notes(slide_data)
    return canvas.preview


def main():
    deck_path = Path(sys.argv[1]).resolve()
    version = sys.argv[2]
    folder = deck_path.parent
    deck = json.loads(deck_path.read_text(encoding='utf-8'))
    colors = palette_for(deck.get('context', {}).get('primary_color'))
    prs = Presentation()
    prs.slide_width = unit(WIDTH)
    prs.slide_height = unit(HEIGHT)
    preview_dir = folder / f'preview-{version}'
    preview_dir.mkdir(parents=True, exist_ok=True)
    for index, slide_data in enumerate(deck['slides']):
        image = render_slide(prs, slide_data, index, len(deck['slides']), folder,
                             deck.get('context', {}), colors)
        image.save(preview_dir / f'{index}.png', optimize=True)
    temporary = folder / f'{version}.tmp.pptx'
    final = folder / f'{version}.pptx'
    prs.save(temporary)
    temporary.replace(final)
    print('OK')


if __name__ == '__main__':
    main()
